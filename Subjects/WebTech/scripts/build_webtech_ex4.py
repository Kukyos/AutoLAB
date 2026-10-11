"""Build Cleo's WebTech Ex4 record -- same pattern as build_webtech_ex1/2/3.py:
fills mam's Lab Record Format.docx template with the real advanced-CSS3 demo
page, the real code, and real screenshots.

    python Subjects/WebTech/build_webtech_ex4.py

index.html + its local assets + the output_*.jpg screenshots must already
exist in output/Ex4/ (written by hand + captured via the browser, not by this
script) -- this script only assembles the docx around them.
"""
from __future__ import annotations

import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from assemble import add_code_block

HERE = Path(__file__).resolve().parents[1]
MATERIALS = HERE / "materials"
OUT = HERE / "output" / "Ex4"
URK = "URK24CS1021"
COURSE = "23CS2049 - Web Technology Lab"
TITLE = "Webpage Design Using Advanced CSS3"
DATE = "24/08/2026"
GITHUB_LINK = "https://kukyos.github.io/WebTechLab/Ex4/index.html"
IMG_W = 5.0

AIM = ("To design a webpage that demonstrates advanced CSS3 styling and "
       "selector features -- rounded corners, border images, multiple "
       "backgrounds, background sizing, pseudo-classes, the adjacent "
       "sibling selector and nth-child() -- applied to the online gaming "
       "community theme.")

DESC = ("Seven CSS3 features are demonstrated as one page, each in its own "
        "themed section. border-radius rounds the corner of four rank "
        "badges, from a single value up to four independent corners and a "
        "50% circle. border-image replaces a plain border with a sliced, "
        "repeated PNG frame on a notice board. Two comma-separated "
        "background-image layers stack a crest emblem (no-repeat, bottom "
        "right) over a repeating stripe pattern on a guild banner. The same "
        "banner image is shown twice at a fixed box size with "
        "background-size: cover and contain, to compare cropping against "
        "letterboxing. A quick-join form uses :focus, :required and "
        "input[type=email]:valid/:invalid to react live to what is typed. "
        "The adjacent sibling combinator (.item-name:hover + .item-img) "
        "fades a loadout icon in only when the item name right before it is "
        "hovered. Finally tr:nth-child(odd/even) stripes a leaderboard "
        "table and a plain :hover rule highlights whichever row the cursor "
        "is over.")

PROCEDURE = (
    "1. Generated the local image assets needed for the demo with PIL "
    "(generate_assets.py): a border-image frame tile, a crest emblem and a "
    "stripe-pattern tile for the layered background, and three small item "
    "icons -- no images were pulled from the network.\n"
    "2. Built the Rank Badges section: four badges with border-radius set "
    "to one value, a different single value, four separate corner values, "
    "and 50% to make a circle.\n"
    "3. Built the Event Notice Board section: a transparent-bordered box "
    "styled with border-image: url(border_deco.png) 20 round.\n"
    "4. Built the Guild Banner section: background-image given two "
    "comma-separated url()s, with background-position and "
    "background-repeat also given one value per layer so the emblem sits "
    "still while the pattern behind it tiles.\n"
    "5. Built the Banner Preview section: the same banner.png shown in two "
    "identically sized boxes, one with background-size: cover and one with "
    "contain, to compare how each handles an image that doesn't match the "
    "box's aspect ratio.\n"
    "6. Built the Quick Join form: text and email inputs styled with "
    ":focus (active field), :required (mandatory field) and "
    "input[type=email]:valid / :invalid (live email-format feedback), "
    "checking the cascade order so :focus visibly wins over :required when "
    "a field is both.\n"
    "7. Built the Loadout section: an item name next to an icon image, "
    "with .item-img given opacity 0.2 and .item-name:hover + .item-img "
    "given opacity 1 plus a transition, so hovering the name reveals the "
    "icon beside it.\n"
    "8. Built the Server Leaderboard table: tr:nth-child(odd) and "
    "tr:nth-child(even) given different background colors to stripe the "
    "rows, plus a tr:hover rule on top for the highlight.\n"
    "9. Served the page locally and tested every feature for real in the "
    "browser -- typed into the join form to see :focus/:invalid react, "
    "hovered a loadout item and a leaderboard row -- before taking the "
    "output screenshots."
)

RESULT = ("The webpage was successfully designed using advanced CSS3 "
          "features. Rounded corners, a border image, layered backgrounds, "
          "background sizing, pseudo-classes, the adjacent sibling "
          "selector and nth-child() all rendered and behaved as expected "
          "in the browser.")


def set_placeholder(doc: Document, marker: str, text: str) -> None:
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip() == marker:
            ph = paras[i + 1]
            for r in list(ph.runs):
                r.text = ""
            if ph.runs:
                ph.runs[0].text = text
            else:
                ph.add_run(text)
            return
    raise ValueError(f"placeholder for {marker!r} not found")


def clear_placeholder(doc: Document, marker: str) -> "docx.text.paragraph.Paragraph":
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip() == marker:
            ph = paras[i + 1]
            for r in list(ph.runs):
                r.text = ""
            return ph
    raise ValueError(f"placeholder for {marker!r} not found")


def fill_header_regno(doc: Document, urk: str) -> None:
    """See build_webtech_ex1.py's copy of this function for why."""
    from docx.oxml.ns import qn
    header = doc.sections[0].header._element
    for t in header.iter(qn("w:t")):
        if t.text == "[regno]":
            t.text = urk
            return
    raise ValueError("[regno] placeholder not found in header")


def fill_footer_title(doc: Document, ex_label: str, title: str) -> None:
    """See build_webtech_ex1.py's copy."""
    from docx.oxml.ns import qn
    footer = doc.sections[0].footer._element
    found_ex, found_title = False, False
    for t in footer.iter(qn("w:t")):
        if t.text == "Ex. 1":
            t.text = ex_label
            found_ex = True
        elif t.text and t.text.strip() == "| Title of exercise":
            t.text = f" | {title}"
            found_title = True
    if not (found_ex and found_title):
        raise ValueError(f"footer placeholders not found (ex={found_ex}, title={found_title})")


def set_ex_number(docx_path: Path, label: str) -> None:
    """See build_webtech_ex1.py's copy: patches docProps/app.xml directly."""
    import re, zipfile
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                text = data.decode("utf8")
                text = re.sub(r"<Company>.*?</Company>", f"<Company>{label}</Company>", text)
                data = text.encode("utf8")
            zout.writestr(item, data)
    tmp.replace(docx_path)


def build() -> Path:
    tmpl = next(MATERIALS.glob("Lab Record Format.docx"))
    doc = Document(str(tmpl))

    doc.core_properties.title = COURSE
    doc.core_properties.author = URK
    doc.core_properties.last_modified_by = URK
    fill_header_regno(doc, URK)
    fill_footer_title(doc, "Ex. 4", TITLE)

    t = doc.tables[0]
    t.rows[0].cells[0].paragraphs[0].runs[-1].text = "4"
    t.rows[0].cells[1].paragraphs[0].runs[0].text = TITLE
    for r in t.rows[0].cells[1].paragraphs[0].runs[1:]:
        r.text = ""
    dcell = t.rows[1].cells[1]
    if not dcell.paragraphs[0].runs:
        dcell.paragraphs[0].add_run(DATE)
    else:
        dcell.paragraphs[0].runs[0].text = DATE
    gcell = t.rows[2].cells[1]
    if not gcell.paragraphs[0].runs:
        gcell.paragraphs[0].add_run(GITHUB_LINK)
    else:
        gcell.paragraphs[0].runs[0].text = GITHUB_LINK

    set_placeholder(doc, "Aim", AIM)
    set_placeholder(doc, "Description", DESC)
    set_placeholder(doc, "Procedure:", PROCEDURE)

    program_anchor = clear_placeholder(doc, "Program")
    add_code_block(doc, (OUT / "index.html").read_text(encoding="utf8"))
    new_block = doc.paragraphs[-1]
    program_anchor._p.addnext(new_block._p)

    output_anchor = clear_placeholder(doc, "Output")
    insert_after = output_anchor
    for fname, caption in [
        ("output_1_top.jpg", "Rank badges (border-radius) and the Event Notice Board (border-image)"),
        ("output_2_backgrounds.jpg", "Guild Banner (layered backgrounds) and Banner Preview (background-size: cover vs contain)"),
        ("output_3_pseudoclasses.jpg", "Quick Join form: :focus (orange border) and a live :invalid email (red text)"),
        ("output_4_leaderboard.jpg", "Loadout icons revealed via adjacent sibling hover, and the striped leaderboard with row 3 hovered"),
    ]:
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption); cr.italic = True; cr.font.size = Pt(10)
        img = doc.add_paragraph(); img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img.paragraph_format.space_after = Pt(10)
        img.add_run().add_picture(str(OUT / fname), width=Inches(IMG_W))
        insert_after._p.addnext(cap._p)
        cap._p.addnext(img._p)
        insert_after = img

    set_placeholder(doc, "Result", RESULT)

    for p in list(doc.paragraphs):
        if p.text.strip().startswith("[Kindly Note"):
            p._p.getparent().remove(p._p)

    out_path = HERE / "records" / "Ex4" / f"Ex4_{URK}.docx"; out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    set_ex_number(out_path, "Ex. 4")
    print("wrote", out_path)
    return out_path


if __name__ == "__main__":
    build()
