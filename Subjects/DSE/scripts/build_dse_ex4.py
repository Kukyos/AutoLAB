"""Build Cleo's DSE Ex-4 record (descriptive statistics + correlation on
fanspeed.csv) from scratch, same approach as build_dse_ex1/2/3.py: run the real
Python, capture real output and real plots, assemble the .docx.

    python Subjects/DSE/build_dse_ex4.py
"""
from __future__ import annotations

import ast, base64, contextlib, io, json, os, shutil, sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")  # headless: this runs from a build script, never a GUI
import matplotlib.pyplot as plt

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from assemble import add_code_block
import labshot

HERE = Path(__file__).resolve().parents[1]
OUT = HERE / "output" / "Ex4"
SRC_CSV = HERE / "output" / "Ex2" / "fanspeed.csv"  # same dataset as Ex-2
URK = "URK24CS1021"
DATE = "07/08/2026"
COURSE = "23CS2025 DATA SCIENCE ECOSYSTEM LAB"
TITLE = "DESCRIPTIVE STATISTICS AND CORRELATION"
FONT = "Times New Roman"

AIM = ("To develop and implement Python programs that summarise real-world datasets "
       "using measures of central tendency, dispersion and correlation, and to "
       "visualise those relationships using Matplotlib and Seaborn.")
DESC = ("Descriptive statistics reduce a column of numbers to a few values that describe "
        "it: the mean, median and mode say where the data sits, while the variance, "
        "standard deviation, range and inter-quartile range say how spread out it is. "
        "Correlation goes one step further and measures whether two columns move "
        "together \u2014 Pearson for a straight-line relationship and Spearman for rank "
        "order. Pandas computes all of these directly on a DataFrame, and Seaborn draws "
        "them as heatmaps, regression plots, box plots, histograms and pair plots.")

QUESTION = (
    "Dataset: fanspeed.csv (https://dataset.karunya.edu/CSE/fanspeed.csv)\n"
    "Columns: Temperature, Observed Fan Speed\n"
    "\u2022 Calculate the mean (average) value for each numerical column in the dataset.\n"
    "\u2022 Determine the median for each numerical column. How does it compare to the mean?\n"
    "\u2022 Identify the mode(s) for categorical and numerical columns. Are there any "
    "columns with multiple modes?\n"
    "\u2022 Compute the variance for each numerical column. Which column has the highest "
    "variability?\n"
    "\u2022 Find the standard deviation for each numerical column. How does it relate to "
    "variance?\n"
    "\u2022 Calculate the range (max - min) for each numerical column. Which column has "
    "the widest range?\n"
    "\u2022 Determine the IQR for each numerical column. Are there any outliers based on "
    "the 1.5\u00d7IQR rule?\n"
    "\u2022 Compute the correlation matrix for all numerical variables. Which pairs have "
    "the strongest correlations?\n"
    "\u2022 Calculate the Pearson correlation coefficient between two selected numerical "
    "columns. Is the relationship positive, negative, or neutral?\n"
    "\u2022 Compute the Spearman rank correlation for two selected columns. How does it "
    "compare to the Pearson correlation?\n"
    "\u2022 Plot a heatmap of the correlation matrix. Which variables are most strongly "
    "correlated?\n"
    "\u2022 Visualize the relationship between two numerical variables with a scatter plot "
    "and regression line. Does the trend match the correlation coefficient?\n"
    "\u2022 Create a box plot for a numerical column. Are there any outliers?\n"
    "\u2022 Plot a histogram for a numerical column. Is the distribution normal, skewed, "
    "or bimodal?\n"
    "\u2022 Generate a pair plot for multiple numerical variables. What patterns do you "
    "observe?")
RESULT = "The exercise shown above has been successfully executed and the output has been verified."

PREAMBLE = (
    "import pandas as pd\n"
    "import matplotlib.pyplot as plt\n"
    "import seaborn as sns\n\n"
    'sns.set_theme(style="whitegrid")\n'
    'plt.rcParams["figure.figsize"] = (7, 4.5)\n'
    'pd.set_option("display.width", 120)\n\n'
)

# (heading for the docx label, source code for that question)
SECTIONS: list[tuple[str, str]] = [
    ("Loading the dataset",
     'df = pd.read_csv("fanspeed.csv")\n'
     'print(df.head())\n'
     'print("\\nShape:", df.shape)\n'
     'print("\\nData types:")\nprint(df.dtypes)'),

    ("1. Mean of each numerical column",
     'print("Mean:")\nprint(df.mean(numeric_only=True))'),

    ("2. Median of each numerical column, compared to the mean",
     'print("Median:")\nprint(df.median(numeric_only=True))\n\n'
     'compare = pd.DataFrame({"mean": df.mean(numeric_only=True),\n'
     '                        "median": df.median(numeric_only=True)})\n'
     'compare["difference"] = compare["mean"] - compare["median"]\n'
     'print("\\nMean vs median:")\nprint(compare)'),

    ("3. Mode(s) for numerical and categorical columns",
     '# every reading in this dataset is unique, so every value ties as a mode\n'
     'for col in df.columns:\n'
     '    modes = df[col].mode()\n'
     '    print(col, "-> number of modes:", len(modes))\n'
     '    print("   first few:", list(modes.head(3)))\n\n'
     '# the dataset has no categorical column, so derive one by binning temperature\n'
     'df["Temp_Band"] = pd.cut(df["Temperature"], bins=[5, 15, 25, 35],\n'
     '                          labels=["Low", "Mid", "High"], include_lowest=True)\n'
     'print("\\nCounts per band:")\nprint(df["Temp_Band"].value_counts())\n'
     'print("\\nMode of the categorical column:", list(df["Temp_Band"].mode()))'),

    ("4. Variance of each numerical column",
     'print("Variance:")\nprint(df.var(numeric_only=True))\n'
     'print("\\nHighest variance:", df.var(numeric_only=True).idxmax())'),

    ("5. Standard deviation, and how it relates to variance",
     'print("Standard deviation:")\nprint(df.std(numeric_only=True))\n\n'
     'check = pd.DataFrame({"variance": df.var(numeric_only=True),\n'
     '                      "std": df.std(numeric_only=True)})\n'
     'check["std squared"] = check["std"] ** 2\n'
     'print("\\nStd is the square root of variance:")\nprint(check)'),

    ("6. Range (max - min) of each numerical column",
     'nums = df.select_dtypes(include="number")\n'
     'rng = pd.DataFrame({"min": nums.min(), "max": nums.max()})\n'
     'rng["range"] = rng["max"] - rng["min"]\n'
     'print(rng)\n'
     'print("\\nWidest range:", rng["range"].idxmax())'),

    ("7. Inter-quartile range and the 1.5xIQR outlier rule",
     'q1 = nums.quantile(0.25)\n'
     'q3 = nums.quantile(0.75)\n'
     'iqr = q3 - q1\n'
     'fences = pd.DataFrame({"Q1": q1, "Q3": q3, "IQR": iqr,\n'
     '                       "lower fence": q1 - 1.5 * iqr,\n'
     '                       "upper fence": q3 + 1.5 * iqr})\n'
     'print(fences)\n\n'
     'print("\\nOutliers per column:")\n'
     'for col in nums.columns:\n'
     '    low = q1[col] - 1.5 * iqr[col]\n'
     '    high = q3[col] + 1.5 * iqr[col]\n'
     '    out = nums[(nums[col] < low) | (nums[col] > high)]\n'
     '    print("  ", col, "->", len(out), "outliers")'),

    ("8. Correlation matrix for all numerical variables",
     'corr = nums.corr()\n'
     'print(corr)\n'
     '# only two numerical columns, so there is exactly one pair to compare\n'
     'print("\\nTemperature vs Observed Fan Speed:",\n'
     '      round(corr.loc["Temperature", "Observed Fan Speed"], 4))'),

    ("9. Pearson correlation between two selected columns",
     'pearson = df["Temperature"].corr(df["Observed Fan Speed"], method="pearson")\n'
     'print("Pearson correlation:", round(pearson, 4))\n'
     'print("Direction:", "positive" if pearson > 0 else\n'
     '      ("negative" if pearson < 0 else "neutral"))'),

    ("10. Spearman rank correlation, compared to Pearson",
     'spearman = df["Temperature"].corr(df["Observed Fan Speed"], method="spearman")\n'
     'print("Spearman correlation:", round(spearman, 4))\n'
     'print("Pearson correlation :", round(pearson, 4))\n'
     'print("Difference          :", round(pearson - spearman, 4))'),

    ("11. Heatmap of the correlation matrix",
     'sns.heatmap(corr, annot=True, fmt=".3f", cmap="coolwarm",\n'
     '            vmin=-1, vmax=1, square=True)\n'
     'plt.title("Correlation matrix")\n'
     'plt.tight_layout()\n'
     'plt.show()'),

    ("12. Scatter plot with a regression line",
     'sns.regplot(data=df, x="Temperature", y="Observed Fan Speed",\n'
     '            scatter_kws={"s": 30, "alpha": 0.7},\n'
     '            line_kws={"color": "red"})\n'
     'plt.title("Observed Fan Speed vs Temperature, with regression line")\n'
     'plt.tight_layout()\n'
     'plt.show()'),

    ("13. Box plot for a numerical column",
     'sns.boxplot(y=df["Observed Fan Speed"], color="#8FBC8F")\n'
     'plt.title("Observed Fan Speed - box plot")\n'
     'plt.tight_layout()\n'
     'plt.show()'),

    ("14. Histogram for a numerical column",
     'print("Skewness:", round(df["Observed Fan Speed"].skew(), 4))\n'
     'sns.histplot(df["Observed Fan Speed"], bins=12, kde=True, color="#4C72B0")\n'
     'plt.title("Distribution of Observed Fan Speed")\n'
     'plt.xlabel("Observed Fan Speed")\n'
     'plt.tight_layout()\n'
     'plt.show()'),

    ("15. Pair plot of the numerical variables",
     'sns.pairplot(nums, diag_kind="hist", height=2.6)\n'
     'plt.suptitle("Pair plot - Temperature and Observed Fan Speed", y=1.02)\n'
     'plt.show()'),
]


def shot(output: str) -> bytes:
    return labshot.render_terminal_shot(
        output, command="python fanspeed_stats.py", cwd=r"C:\DSE\Ex4", style="cmd",
        urk_mode="none", show_banner=False, trailing_prompt=True, width=1500)


def _grab_figure() -> bytes | None:
    """Save whatever figure was just drawn, then close it. Figure-level seaborn
    calls (pairplot) make their own figure, which is the one gcf() returns."""
    if not plt.get_fignums():
        return None
    buf = io.BytesIO()
    plt.gcf().savefig(buf, format="png", dpi=130, bbox_inches="tight")
    plt.close("all")
    return buf.getvalue()


def run_sections() -> list[list[bytes]]:
    """Exec each section cumulatively in one shared namespace. Each section gives
    back its images: a terminal shot if it printed anything, plus its plot."""
    ns: dict = {}
    prev_cwd = os.getcwd()
    os.chdir(OUT)  # fanspeed_stats.py reads "fanspeed.csv" next to itself
    try:
        exec(compile(PREAMBLE, "<preamble>", "exec"), ns)
        shots = []
        for heading, code in SECTIONS:
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                exec(compile(code, "<section>", "exec"), ns)
            imgs = []
            if buf.getvalue().strip():
                imgs.append(shot(buf.getvalue().rstrip("\n")))
            fig = _grab_figure()
            if fig:
                imgs.append(fig)
            assert imgs, f"section produced no output: {heading}"
            shots.append(imgs)
    finally:
        os.chdir(prev_cwd)
    return shots


# ---- notebook + video script ---------------------------------------------
NB_CELLS = [
    ("md", "# Ex 4 \u2014 Descriptive Statistics and Correlation\n\n"
           "**Dataset:** fanspeed.csv  \n**Reg No:** URK24CS1021", None),
    ("code",
     "import pandas as pd\n"
     "import matplotlib.pyplot as plt\n"
     "import seaborn as sns\n\n"
     'sns.set_theme(style="whitegrid")\n'
     'plt.rcParams["figure.figsize"] = (7, 4.5)\n\n'
     'df = pd.read_csv("fanspeed.csv")\n'
     "df.head()", None),
    ("md", "### Where does the data sit? \u2014 mean, median, mode", None),
    ("code",
     'pd.DataFrame({"mean": df.mean(), "median": df.median()})', None),
    ("code",
     "# every reading is unique, so every value ties as a mode\n"
     'len(df["Temperature"].mode()), len(df["Observed Fan Speed"].mode())', None),
    ("code",
     "# no categorical column in the file, so bin temperature into one\n"
     'df["Temp_Band"] = pd.cut(df["Temperature"], bins=[5, 15, 25, 35],\n'
     '                         labels=["Low", "Mid", "High"], include_lowest=True)\n'
     'df["Temp_Band"].mode()', None),
    ("md", "### How spread out is it? \u2014 variance, std, range, IQR", None),
    ("code",
     "nums = df.select_dtypes(include=\"number\")\n"
     'pd.DataFrame({"variance": nums.var(), "std": nums.std(),\n'
     '              "range": nums.max() - nums.min()})', None),
    ("code",
     "q1, q3 = nums.quantile(0.25), nums.quantile(0.75)\n"
     "iqr = q3 - q1\n"
     'pd.DataFrame({"IQR": iqr, "lower fence": q1 - 1.5 * iqr,\n'
     '              "upper fence": q3 + 1.5 * iqr})', None),
    ("code",
     "# how many rows fall outside the 1.5xIQR fences\n"
     'col = "Observed Fan Speed"\n'
     "low, high = q1[col] - 1.5 * iqr[col], q3[col] + 1.5 * iqr[col]\n"
     "((nums[col] < low) | (nums[col] > high)).sum()", None),
    ("md", "### Do the two columns move together? \u2014 correlation", None),
    ("code", "corr = nums.corr()\ncorr", None),
    ("code",
     'pearson = df["Temperature"].corr(df["Observed Fan Speed"], method="pearson")\n'
     'spearman = df["Temperature"].corr(df["Observed Fan Speed"], method="spearman")\n'
     "round(pearson, 4), round(spearman, 4)", None),
    ("code",
     'sns.heatmap(corr, annot=True, fmt=".3f", cmap="coolwarm", vmin=-1, vmax=1, square=True)\n'
     'plt.title("Correlation matrix")\n'
     "plt.show()", None),
    ("md", "### Seeing it \u2014 regression line, box plot, histogram, pair plot", None),
    ("code",
     'sns.regplot(data=df, x="Temperature", y="Observed Fan Speed", line_kws={"color": "red"})\n'
     'plt.title("Fan speed vs temperature")\n'
     "plt.show()", None),
    ("code",
     'sns.boxplot(y=df["Observed Fan Speed"], color="#8FBC8F")\n'
     'plt.title("Observed Fan Speed - box plot")\n'
     "plt.show()", None),
    ("code",
     'df["Observed Fan Speed"].skew()', None),
    ("code",
     'sns.histplot(df["Observed Fan Speed"], bins=12, kde=True)\n'
     'plt.title("Distribution of Observed Fan Speed")\n'
     "plt.show()", None),
    ("code",
     'sns.pairplot(nums, diag_kind="hist", height=2.6)\n'
     "plt.show()", None),
]


def _srclist(text: str) -> list[str]:
    return text.splitlines(keepends=True) or [""]


def _run_cell(src: str, ns: dict, count: int) -> list[dict]:
    tree = ast.parse(src)
    last = tree.body[-1] if tree.body else None
    buf, result = io.StringIO(), None
    with contextlib.redirect_stdout(buf):
        if isinstance(last, ast.Expr):
            exec(compile(ast.Module(tree.body[:-1], []), "<cell>", "exec"), ns)
            result = eval(compile(ast.Expression(last.value), "<cell>", "eval"), ns)
        else:
            exec(compile(tree, "<cell>", "exec"), ns)
    outs = []
    if buf.getvalue():
        outs.append({"output_type": "stream", "name": "stdout",
                     "text": _srclist(buf.getvalue())})
    fig = _grab_figure()  # plot cells return None from plt.show(), so catch the figure
    if fig:
        outs.append({"output_type": "display_data", "metadata": {},
                     "data": {"image/png": base64.b64encode(fig).decode()}})
    if result is not None:
        data = {"text/plain": _srclist(repr(result))}
        if hasattr(result, "_repr_html_"):
            data["text/html"] = _srclist(result._repr_html_())
        outs.append({"output_type": "execute_result", "execution_count": count,
                     "metadata": {}, "data": data})
    return outs


def build_notebook(path: Path) -> None:
    ns: dict = {}
    cells, count = [], 0
    for kind, src, _note in NB_CELLS:
        if kind == "md":
            cells.append({"cell_type": "markdown", "metadata": {}, "source": _srclist(src)})
        else:
            count += 1
            # teacher's rule from Ex-2 onward: reg no printed first, every cell
            src = f'print("{URK}")\n\n' + src
            cells.append({"cell_type": "code", "execution_count": count,
                          "metadata": {}, "outputs": _run_cell(src, ns, count),
                          "source": _srclist(src)})
    nb = {"cells": cells,
          "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python",
                                      "name": "python3"},
                       "language_info": {"name": "python", "version": "3.14"}},
          "nbformat": 4, "nbformat_minor": 5}
    path.write_text(json.dumps(nb, indent=1), encoding="utf8")


VIDEO = """# Ex 4 — Code explanation (teaching cues)
*Descriptive statistics & correlation · fanspeed.csv · URK24CS1021*

> Mam wants this to sound like **you teaching a junior**, not you reading the code.
> So the order is always: **the question first, why this measure and not another, then the code.**
> Glance at these cues, look at the output on screen, say it in your own words. ~6-7 min.

**Open with** — who you are, reg no URK24CS1021, Experiment 4.
Then the hook, in your own words: *"Same fan dataset as Experiment 2 — 100 readings,
temperature and the fan speed observed at it. Last time I looked at the rows. This time
I never look at a single row. The whole idea of statistics is to replace a hundred numbers
with about six that tell you the same story: where the data sits, how spread out it is,
and whether the two columns move together. Then I plot it, because a number can lie about
a shape and a picture can't."*

---

**Where the data sits — mean, median, mode (Q1–Q3)**
- Question: *what's a typical reading?*
- Do mean and median together, because the interesting thing isn't either number — it's
  **the gap between them**. They're almost identical here, and that tells you the data is
  roughly symmetric. If the mean were far above the median you'd know a few big values
  were dragging it up. That's the entire reason you compute both.
- Why the median resists that: it only cares about the middle position, so one enormous
  outlier moves it by one place. The mean gets pulled the whole way.
- Then the honest bit on mode: **every one of the 100 readings is unique**, so `mode()`
  comes back with all 100 values tied. That's not a bug — it's what mode means on
  continuous data, and the lesson is that **mode is a categorical tool.** So I bin
  temperature into Low/Mid/High with `pd.cut` and take the mode of *that*, which is the
  answer the question is actually after.

**How spread out it is — variance, std, range, IQR (Q4–Q7)**
- Question: *are the readings tightly packed or all over the place?*
- Variance first: average squared distance from the mean. Ask out loud why it's *squared* —
  because plain distances cancel out to zero, half above and half below.
- Then the payoff: **standard deviation is just the square root of variance**, and the
  reason you report std instead is **units**. Variance of fan speed is in "fan-speed
  squared", which means nothing. Std is in fan speed. Point at the check table where
  std² lands back on the variance — that's the relationship, shown rather than claimed.
- Range is max − min: one line, and one weakness — it's decided entirely by the two most
  extreme readings, so a single bad sensor value defines it.
- **IQR fixes exactly that.** It's the range of the middle 50%, so it ignores the tails.
  Then the 1.5×IQR rule: anything more than one and a half IQRs past the quartiles is
  flagged as an outlier. Read out the fence numbers and the count you actually got —
  don't claim a number you haven't looked at.
- Fan speed has a much bigger variance than temperature. Say why that's expected before
  someone asks: the two columns are on different scales, so comparing their spreads
  directly isn't very meaningful. That's a real caveat, not a footnote.

**Do the columns move together — correlation (Q8–Q10)**
- Question: *when temperature goes up, does fan speed follow?*
- Explain the scale once: −1 is a perfect downward line, 0 is no linear relationship,
  +1 is a perfect upward line. Ours is about 0.985 — very strong and positive.
- **Pearson vs Spearman is the concept worth the time.** Pearson measures how close the
  points are to a *straight line*. Spearman throws away the values and keeps only the
  *rank order*, then measures that. Here they're almost identical, which tells you the
  relationship is both increasing *and* basically straight.
- Say when they'd disagree, because that's what makes the pair useful: a curve that keeps
  going up — doubling every step, say — scores near 1 on Spearman and noticeably lower on
  Pearson. If you ever see those two numbers far apart, that's your signal the
  relationship is real but not a straight line.
- The caveat, said once and said plainly: **correlation is not causation.** These two
  really are causally linked, but the number alone could never prove that.

**Seeing it — heatmap, regression, box, histogram, pair plot (Q11–Q15)**
- Frame the section: *the numbers said 0.985. Now let's check the picture agrees.*
- **Heatmap:** it colours the correlation matrix so strong pairs jump out. Be upfront —
  with only two numerical columns this matrix is 2×2 and the diagonal is 1.0 by
  definition, so there's exactly one real number in it. On a dataset with twenty columns
  this is the plot that saves you reading four hundred numbers.
- **Regplot:** scatter plus the fitted line in one call. Point at it: the points hug the
  line, which is what a 0.985 is *supposed* to look like. That's the check — a
  correlation number with no plot behind it is a claim you haven't verified.
- **Box plot:** ties straight back to the IQR you computed. The box is Q1 to Q3, the line
  is the median, the whiskers stop at the 1.5×IQR fences, and any dot beyond them is an
  outlier. Say whether you see any, and match it against the count you printed earlier.
- **Histogram:** the shape question. Don't eyeball it — I printed the skewness, so read
  the number and then say what it means: near zero is symmetric, positive means a tail to
  the right, negative to the left. The KDE curve on top is a smoothed version of the same
  bars.
- **Pair plot:** every numerical column against every other, with each column's own
  distribution down the diagonal. Say what it's *for* — on a wide dataset it's the first
  thing you run, one glance and you know which pairs are worth investigating. With two
  columns here it's a 2×2 grid, so it's the technique on show, not the finding.

---

**Close with** — the workflow, not the function names:
*where does it sit (mean/median/mode) → how spread out is it (variance/std/range/IQR) →
do the columns move together (Pearson/Spearman) → and does a plot agree with the numbers.*
Then the one line to leave them with: every summary statistic throws information away on
purpose — that's why you always plot it before you trust it.
"""


# ---- docx helpers (same as build_dse_ex1/2/3.py) --------------------------
def rfonts(el, name=FONT):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)


def cell_border(cell):
    pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def label(doc, text, size=12, space_before=8):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.name = FONT; r.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    return p


def body(doc, text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(12); r.italic = italic
    p.paragraph_format.space_after = Pt(4)
    return p


def _img(doc, png, width_in=6.2):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "screenshots").mkdir(exist_ok=True)
    shutil.copyfile(SRC_CSV, OUT / "fanspeed.csv")  # same file as Ex-2, not re-fetched

    full_code = PREAMBLE + "\n\n".join(
        f"# {heading}\n{code}" for heading, code in SECTIONS)
    (OUT / "fanspeed_stats.py").write_text(
        "# Ex 4 - Descriptive Statistics and Correlation (fanspeed.csv)\n"
        "# URK24CS1021\n\n" + full_code, encoding="utf8")

    shots = run_sections()
    for i, imgs in enumerate(shots, 1):
        for j, png in enumerate(imgs):
            suffix = "" if len(imgs) == 1 else chr(ord("a") + j)
            (OUT / "screenshots" / f"q{i:02d}{suffix}_output.png").write_bytes(png)

    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = FONT; normal.font.size = Pt(12)
    rfonts(normal.element)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))

    h = sec.header.paragraphs[0]
    h.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    r = h.add_run(COURSE); r.bold = True
    h.add_run("\t"); r2 = h.add_run(URK); r2.bold = True
    pbdr = h._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "000000")
    bd.append(bot); pbdr.append(bd)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = f.add_run(f"EX-4: {TITLE}"); fr.bold = True

    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    rows = [("Ex. No. 4", TITLE), ("Date of Exercise", DATE)]
    widths = (Inches(1.9), Inches(4.35))
    for ri, (a, b) in enumerate(rows):
        for ci, txt in enumerate((a, b)):
            c = t.rows[ri].cells[ci]; c.width = widths[ci]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = p.add_run(txt); cr.font.name = FONT; cr.font.size = Pt(12)
            cr.bold = (ci == 0 or ri == 0)
            cell_border(c)
    doc.add_paragraph()

    label(doc, "Aim"); body(doc, AIM)
    label(doc, "Description (About Descriptive Statistics & Correlation)"); body(doc, DESC)
    label(doc, "Question:"); body(doc, QUESTION)

    for (heading, code), imgs in zip(SECTIONS, shots):
        label(doc, f"Sample Code: {heading}")
        add_code_block(doc, code)
        label(doc, "Sample Output:")
        for png in imgs:
            _img(doc, png)

    label(doc, "CODE EXPLANATION LINK:")
    body(doc, "(Youtube / Drive Link)", italic=True)
    label(doc, "Preparation course certificate")
    doc.add_paragraph()
    label(doc, "RESULT"); body(doc, RESULT)

    out_docx = HERE / "records" / "Ex4" / f"Ex4_{URK}.docx"; out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))

    prev_cwd = os.getcwd()
    os.chdir(OUT)
    try:
        build_notebook(OUT / "Ex4_FanSpeedStats.ipynb")
    finally:
        os.chdir(prev_cwd)
    (OUT / "video_script.md").write_text(VIDEO, encoding="utf8")
    print("wrote", OUT)
    print("  sections:", len(shots), "| images:", sum(len(i) for i in shots),
          "| docx:", out_docx.name, "| notebook + video_script.md")


if __name__ == "__main__":
    build()
