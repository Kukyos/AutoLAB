"""Build a lab record .docx from scratch, from a JSON spec.

For when the teacher gave no editable template (only a question PDF, say).
The layout is the one that held up across real submissions: a header with
course on the left and register number on the right over a rule, an
"Ex. No. / Date" table, Aim, Description, then each question with its code
and its output screenshots, and a Result section. Footer: "EX N: TITLE".

Spec (paths are relative to the spec file's folder):

    {
      "course": "23CS2014 Database Systems Lab",
      "register_number": "URK24CS1021",
      "ex_no": "2", "title": "Basic SQL", "date": "29/07/2026",
      "aim": "...", "description": "...",
      "questions": [
        {"number": "1", "text": "Create the table ...",
         "code": "..." | "code_file": "output/Ex2/q1.sql",
         "images": ["output/Ex2/screenshots/q1.png"]}
      ],
      "result": "The queries were executed and the output was verified.",

      // optional
      "font": "Times New Roman", "code_label": "Sample Code:",
      "output_label": "Sample Output:", "image_dpi": 150, "image_max_width_in": 6.3
    }
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image

from .assemble import add_code_block

FONT = "Times New Roman"
SHOT_DPI = 150
IMG_MAX_W = 6.3


def img_width(png: bytes, cap: float = IMG_MAX_W, dpi: int | None = SHOT_DPI) -> Inches:
    """Width for a screenshot: its natural size at `dpi`, never wider than `cap`.

    Placing every shot at one scale keeps console text the same size across
    the record. A fixed width would blow short, narrow shots up until they are
    huge and blurry next to the wide ones. `dpi=None` means always use `cap`.
    """
    if dpi is None:
        return Inches(cap)
    return Inches(min(Image.open(io.BytesIO(png)).width / dpi, cap))


def _rfonts(el, name: str) -> None:
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)


def _cell_border(cell) -> None:
    pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def _label(doc, text: str, font: str, size: int = 12, space_before: int = 8):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.name = font; r.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    return p


def _body(doc, text: str, font: str):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = font; r.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p


def _image(doc, png: bytes, cap: float, dpi: int | None) -> None:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(io.BytesIO(png), width=img_width(png, cap, dpi))


def _header_footer(sec, course: str, reg: str, footer: str) -> None:
    h = sec.header.paragraphs[0]
    h.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    h.add_run(course).bold = True
    h.add_run("\t")
    h.add_run(reg).bold = True
    ppr = h._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "000000")
    bd.append(bot); ppr.append(bd)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.add_run(footer).bold = True


def _resolve(base: Path, rel: str) -> Path:
    p = Path(rel)
    return p if p.is_absolute() else base / p


def build_record(spec: dict, out_path: Path, base_dir: Path | None = None) -> dict:
    """Write the record described by `spec` to `out_path`. Returns a summary."""
    base = base_dir or Path.cwd()
    font = spec.get("font", FONT)
    dpi = spec.get("image_dpi", SHOT_DPI)
    cap = float(spec.get("image_max_width_in", IMG_MAX_W))
    title = spec["title"]
    ex_no = str(spec.get("ex_no", ""))

    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = font; normal.font.size = Pt(12)
    _rfonts(normal.element, font)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))
    _header_footer(sec, spec.get("course", ""), spec.get("register_number", ""),
                   f"EX {ex_no}: {title.upper()}" if ex_no else title.upper())

    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    rows = [(f"Ex. No. {ex_no}", title.upper()), ("Date of Exercise", spec.get("date", ""))]
    widths = (Inches(1.9), Inches(4.35))
    for ri, pair in enumerate(rows):
        for ci, txt in enumerate(pair):
            c = t.rows[ri].cells[ci]; c.width = widths[ci]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt); r.font.name = font; r.font.size = Pt(12)
            r.bold = ci == 0 or ri == 0
            _cell_border(c)
    doc.add_paragraph()

    for key, heading in (("aim", "Aim"), ("description", "Description")):
        if spec.get(key):
            _label(doc, heading, font)
            _body(doc, spec[key], font)

    images_added = 0
    questions = spec.get("questions", [])
    for q in questions:
        n = str(q.get("number", "")).strip()
        _label(doc, f"{n}. {q['text']}" if n else q["text"], font)
        code = q.get("code")
        if code is None and q.get("code_file"):
            code = _resolve(base, q["code_file"]).read_text(encoding="utf-8")
        if code and code.strip():
            if spec.get("code_label", "Sample Code:"):
                _label(doc, spec.get("code_label", "Sample Code:"), font, space_before=2)
            add_code_block(doc, code.rstrip("\n"))
        imgs = q.get("images", [])
        if imgs and spec.get("output_label", "Sample Output:"):
            _label(doc, spec.get("output_label", "Sample Output:"), font, space_before=2)
        for rel in imgs:
            _image(doc, _resolve(base, rel).read_bytes(), cap, dpi)
            images_added += 1

    if spec.get("result"):
        _label(doc, "Result", font)
        _body(doc, spec["result"], font)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return {"out": str(out_path), "questions": len(questions), "images": images_added}
